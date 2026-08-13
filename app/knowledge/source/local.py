"""Local filesystem source adapter.

Reads enterprise knowledge files (PDF / DOCX / TXT) from a root directory
into platform ``SourceDocument`` objects.  It never chunks, embeds, or touches
Qdrant/Haystack.
"""

import os

from app.knowledge.models.document import SourceDocument
from app.knowledge.source.base import KnowledgeSourcePort
from app.knowledge.source.errors import (
    KnowledgeSourceNotFoundError,
    KnowledgeSourceReadError,
)
from app.knowledge.source.scanner import KnowledgeSourceScanner


def _extract_pdf(path):
    """Return list of (page_number, text)."""
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").replace("\x0c", "").strip()
        if text:
            pages.append((index, text))
    return pages


def _extract_docx(path):
    """Return list of (section_heading, text)."""
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    sections = []
    heading = ""
    paragraphs = []
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            if paragraphs:
                sections.append((heading, "\n".join(paragraphs)))
            heading = paragraph.text.strip()
            paragraphs = []
        elif paragraph.text.strip():
            paragraphs.append(paragraph.text.strip())
    if paragraphs:
        sections.append((heading, "\n".join(paragraphs)))
    return sections


def _extract_txt(path):
    with open(path, "r", encoding="utf-8") as handle:
        text = handle.read().strip()
    return [(None, text)] if text else []


def _extract(path, file_type):
    if file_type == "pdf":
        return _extract_pdf(path)
    if file_type == "docx":
        return _extract_docx(path)
    if file_type == "txt":
        return _extract_txt(path)
    raise KnowledgeSourceReadError(f"unsupported file type: {file_type}")


class LocalFileSourceAdapter(KnowledgeSourcePort):
    def __init__(self, root: str):
        if not root:
            raise ValueError("source root is required")
        self.root = root
        self.scanner = KnowledgeSourceScanner(root)

    def list_documents(self):
        return self.scanner.scan()

    def _find(self, source_id: str):
        for document in self.list_documents():
            if document.source_id == source_id:
                return document
        raise KnowledgeSourceNotFoundError(f"source not found: {source_id}")

    def _resolve_path(self, source_id: str) -> str:
        path = os.path.join(self.root, source_id)
        if not os.path.isfile(path):
            raise KnowledgeSourceNotFoundError(f"source not found: {source_id}")
        return path

    def read_document(self, source_id: str) -> SourceDocument:
        document = self._find(source_id)
        path = self._resolve_path(source_id)
        try:
            parts = _extract(path, document.file_type)
        except Exception as error:
            raise KnowledgeSourceReadError(
                f"failed to read source {source_id}: {error}"
            ) from error
        content = "\n\n".join(text for _, text in parts)
        return self._to_source_document(document, content)

    def read_parts(self, source_id: str) -> list[SourceDocument]:
        """Return one SourceDocument per page (PDF) / section (DOCX)."""
        document = self._find(source_id)
        path = self._resolve_path(source_id)
        try:
            parts = _extract(path, document.file_type)
        except Exception as error:
            raise KnowledgeSourceReadError(
                f"failed to read source {source_id}: {error}"
            ) from error
        documents = []
        for index, (marker, text) in enumerate(parts):
            policy = self._access_policy(document)
            if document.file_type == "pdf" and marker is not None:
                policy["page"] = str(marker)
            elif document.file_type == "docx" and marker:
                policy["section"] = marker
            documents.append(
                self._to_source_document(document, text, policy, index)
            )
        return documents

    @staticmethod
    def _access_policy(document) -> dict:
        metadata = document.metadata or {}
        return {
            "security_level": metadata.get("security_level", "public"),
            "region": metadata.get("region"),
            "store_id": metadata.get("store_id"),
            "knowledge_scope": metadata.get("knowledge_scope"),
            "source_id": document.source_id,
            "source_name": document.file_name,
        }

    @staticmethod
    def _to_source_document(document, content, policy=None, index=None) -> SourceDocument:
        policy = policy or LocalFileSourceAdapter._access_policy(document)
        suffix = f":part{index}" if index is not None else ""
        return SourceDocument(
            source_system="local_file",
            external_id=f"{document.source_id}{suffix}",
            tenant_id=document.tenant_id,
            document_type=document.document_type,
            title=document.file_name,
            content=content,
            source_version=document.checksum[:12],
            language="zh",
            access_policy=policy,
        )
