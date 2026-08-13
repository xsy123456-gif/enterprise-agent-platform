"""PDF/DOCX end-to-end ingestion + retrieval evaluation (real files).

Run manually::

    python -m tests.knowledge.evaluation.run_document_evaluation

Validates: file -> convert -> chunk -> embed -> Qdrant -> retrieve, with
per-page/per-section citation, ACL isolation, and hybrid retrieval.
"""

import asyncio
import json
from pathlib import Path

from app.integrations.knowledge.haystack import (
    HaystackIngestionAdapter,
    HaystackKnowledgeConfig,
    HaystackRetrieverAdapter,
    QdrantStoreManager,
)
from app.knowledge import (
    KnowledgeAccessContext,
    KnowledgeConfig,
    KnowledgeRetrieveRequest,
    KnowledgeService,
)
from app.knowledge.ingestion.service import KnowledgeIngestionService
from app.knowledge.models.document import SourceDocument

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"
COLLECTION = "knowledge_doc_eval"

# (filename, source_system, external_id, document_type, security_level, region, store_id)
FIXTURES = [
    ("商品说明书.docx", "product_pim", "SKU-CHARGER", "product_knowledge", "public", "JP", "JP01"),
    ("日本站退换货政策.pdf", "policy_cms", "POLICY-JP-RETURN", "policy", "public", "JP", "JP01"),
    ("广告投放SOP.docx", "ops_sop", "SOP-ADS", "sop", "internal", "JP", "JP01"),
    ("客服FAQ.pdf", "faq_cms", "FAQ-JP", "faq", "public", "JP", "JP01"),
    ("JP01内部规则.pdf", "internal_policy", "JP01-INTERNAL", "internal_policy", "internal", "JP", "JP01"),
    ("管理层制度.docx", "management_policy", "MGMT-POLICY", "management_policy", "confidential", "JP", "JP01"),
]


def _extract_pdf(path):
    """Return list of (page_number, text)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").replace("\x0c", "").strip()
        if text:
            pages.append((index, text))
    return pages


def _extract_docx(path):
    """Return list of (section_heading, text)."""
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    sections = []
    current_heading = ""
    current_paras = []
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            if current_paras:
                sections.append((current_heading, "\n".join(current_paras)))
            current_heading = paragraph.text.strip()
            current_paras = []
        elif paragraph.text.strip():
            current_paras.append(paragraph.text.strip())
    if current_paras:
        sections.append((current_heading, "\n".join(current_paras)))
    return sections


def _to_source_document(fixture, chunk_index, content, page=None, section=None):
    filename, source_system, external_id, doc_type, security, region, store = fixture
    policy = {
        "security_level": security,
        "region": region,
        "store_id": store,
        "source_name": filename,
    }
    if page is not None:
        policy["page"] = str(page)
    if section is not None:
        policy["section"] = section
    return SourceDocument(
        source_system=source_system,
        external_id=f"{external_id}:chunk{chunk_index}",
        tenant_id="tenant_A",
        document_type=doc_type,
        title=filename,
        content=content,
        source_version="2026.08",
        language="zh",
        access_policy=policy,
    )


def _build_documents():
    documents = []
    for fixture in FIXTURES:
        filename = fixture[0]
        path = FIXTURES_DIR / filename
        if filename.endswith(".pdf"):
            for chunk_index, (page, text) in enumerate(_extract_pdf(path)):
                documents.append(_to_source_document(fixture, chunk_index, text, page=page))
        else:
            for chunk_index, (section, text) in enumerate(_extract_docx(path)):
                documents.append(_to_source_document(fixture, chunk_index, text, section=section))
    return documents


def _service(config):
    store_manager = QdrantStoreManager(config)
    ingestion = HaystackIngestionAdapter(config=config, store_manager=store_manager)
    retriever = HaystackRetrieverAdapter(config=config, store_manager=store_manager)
    return (
        KnowledgeIngestionService(ingestion),
        KnowledgeService(retriever, config=KnowledgeConfig()),
        store_manager,
    )


def _ctx(clearance):
    return KnowledgeAccessContext(
        tenant_id="tenant_A", user_id="u1", role="sales",
        authorized_store_ids=("JP01",), authorized_regions=("JP",),
        authorized_knowledge_scopes=(
            "product_knowledge", "policy", "faq", "sop", "internal_policy",
            "management_policy",
        ),
        security_clearance=clearance,
    )


async def main():
    config = HaystackKnowledgeConfig(
        collection=COLLECTION,
        ollama_endpoint="http://172.25.192.1:11434",
        ollama_model="bge-m3",
        embedding_dim=1024,
    )
    ingest_service, service, store_manager = _service(config)

    documents = _build_documents()
    ingested = 0
    for document in documents:
        await ingest_service.ingest(document)
        ingested += 1

    report = {
        "ingested_chunks": ingested,
        "chunks_in_qdrant": store_manager.store.count_documents(),
    }

    def sources(items):
        return {item.citation.source_name for item in items}

    # 1. 正文 + hybrid 检索（public）+ DOCX 章节 citation
    public = _ctx("public")
    r1 = await service.retrieve(KnowledgeRetrieveRequest(query="充电器支持100V吗"), public)
    report["product_retrieval"] = {
        "hit": "商品说明书.docx" in sources(r1.items),
        "section": r1.items[0].citation.section if r1.items else None,
    }

    # 2. PDF 分页 citation
    r2 = await service.retrieve(KnowledgeRetrieveRequest(query="退货期限多少天"), public)
    report["pdf_citation"] = {
        "hit": "日本站退换货政策.pdf" in sources(r2.items),
        "page": r2.items[0].citation.page if r2.items else None,
    }

    # 3. ACL：confidential 管理层制度 不被 public 检索
    r3 = await service.retrieve(KnowledgeRetrieveRequest(query="管理层薪酬制度"), public)
    report["confidential_public_blocked"] = "管理层制度.docx" not in sources(r3.items)

    # 4. ACL：confidential 用户能检索管理层制度
    mgmt = _ctx("confidential")
    r4 = await service.retrieve(KnowledgeRetrieveRequest(query="管理层薪酬制度"), mgmt)
    report["confidential_mgmt_hit"] = "管理层制度.docx" in sources(r4.items)

    # 5. internal SOP 不被 public 检索
    r5 = await service.retrieve(KnowledgeRetrieveRequest(query="广告ACOS升高先查什么"), public)
    report["internal_public_blocked"] = "广告投放SOP.docx" not in sources(r5.items)

    # 6. internal 用户能检索 SOP
    internal = _ctx("internal")
    r6 = await service.retrieve(KnowledgeRetrieveRequest(query="广告ACOS升高先查什么"), internal)
    report["internal_hit"] = "广告投放SOP.docx" in sources(r6.items)

    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    asyncio.run(main())
