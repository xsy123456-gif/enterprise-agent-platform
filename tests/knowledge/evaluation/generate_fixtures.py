"""Generate PDF/DOCX evaluation fixtures (one-time, not part of runtime).

Produces real Chinese e-commerce documents under ``fixtures/`` used by
``run_document_evaluation`` to validate the full file -> convert -> chunk ->
embed -> retrieve pipeline.

PDFs are generated with reportlab's built-in CJK CID font (STSong-Light) so
the text is extractable by standard PDF text extraction.
"""

import os

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

FIXTURES_DIR = os.path.dirname(os.path.abspath(__file__)) + "/fixtures"

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
_TITLE_STYLE = ParagraphStyle(
    "cn-title", fontName="STSong-Light", fontSize=16, leading=20, spaceAfter=12
)
_BODY_STYLE = ParagraphStyle(
    "cn-body", fontName="STSong-Light", fontSize=11, leading=16, spaceAfter=6
)


def _write_pdf(name, pages):
    """pages: list of (title, [paragraphs]).  One chapter per PDF page."""
    doc = SimpleDocTemplate(
        f"{FIXTURES_DIR}/{name}", pagesize=A4,
        title=name, author="evaluation",
    )
    story = []
    for title, paragraphs in pages:
        story.append(Paragraph(title, _TITLE_STYLE))
        for paragraph in paragraphs:
            story.append(Paragraph(paragraph, _BODY_STYLE))
        story.append(PageBreak())
    doc.build(story)
    print("generated", name)


def _write_docx(name, sections):
    """sections: list of (heading, [paragraphs])."""
    doc = DocxDocument()
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
    doc.save(f"{FIXTURES_DIR}/{name}")
    print("generated", name)


def main():
    os.makedirs(FIXTURES_DIR, exist_ok=True)

    # 1. 商品说明书.docx
    _write_docx("商品说明书.docx", [
        ("产品参数", [
            "这款充电器支持100V到240V宽电压输入，全球通用。",
            "输出功率最大65W，兼容主流手机和笔记本电脑。",
        ]),
        ("使用方法", [
            "将充电器接入电源，再连接设备即可开始充电。",
            "长时间不使用时请拔下电源插头。",
        ]),
        ("注意事项", [
            "请勿在潮湿环境下使用充电器。",
            "电池循环寿命约500次，正常使用约两年。",
        ]),
    ])

    # 2. 日本站退换货政策.pdf
    _write_pdf("日本站退换货政策.pdf", [
        ("退货政策", [
            "日本站商品退货期限为30天，自签收之日起计算。",
            "超过30天不支持退货，特殊商品除外。",
        ]),
        ("换货政策", [
            "换货期限为15天，商品须保持未使用状态并保留原包装。",
            "尺寸不合适的商品可以免费申请换货一次。",
        ]),
        ("退款政策", [
            "退款将在7个工作日内原路退回。",
            "跨境订单的退款到账时间可能延长。",
        ]),
    ])

    # 3. 广告投放SOP.docx
    _write_docx("广告投放SOP.docx", [
        ("竞价检查", [
            "广告ACOS突然升高，先检查竞价设置是否被提高。",
            "核对关键词匹配方式和广告预算是否异常。",
        ]),
        ("素材规范", [
            "广告素材不得含有绝对化用语，如最、第一、顶级等。",
        ]),
    ])

    # 4. 客服FAQ.pdf
    _write_pdf("客服FAQ.pdf", [
        ("配送问题", [
            "标准配送时效为3到5个工作日。",
            "偏远地区配送时间会相应延长。",
        ]),
        ("支付问题", [
            "支持信用卡、借记卡和第三方支付。",
            "跨境订单收取一定手续费。",
        ]),
        ("优惠券使用", [
            "优惠券在结算页面输入券码使用，部分商品不参与活动。",
        ]),
    ])

    # 5. JP01内部规则.pdf
    _write_pdf("JP01内部规则.pdf", [
        ("运营规则", [
            "JP01店铺运营人员必须每日核对库存数据。",
            "商品价格调整需经店长审批。",
        ]),
        ("数据安全", [
            "内部经营数据不得外传，违者按公司制度处理。",
        ]),
    ])

    # 6. 管理层制度.docx
    _write_docx("管理层制度.docx", [
        ("薪酬制度", [
            "管理层薪酬结构为基本工资加绩效奖金。",
            "绩效奖金根据季度目标完成情况发放。",
        ]),
        ("审批权限", [
            "单笔金额超过50万元的支出需董事会审批。",
        ]),
    ])


if __name__ == "__main__":
    main()
